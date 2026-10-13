"""Text extraction for every supported upload type. Returns page/slide-aware Segments."""
import csv
import io
from pathlib import PurePath

from app.services.chunking import Segment
from app.services.llm import LLMClient

TEXT_EXT = {".txt", ".md", ".markdown"}
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp"}
AUDIO_EXT = {".mp3", ".wav", ".m4a", ".ogg", ".flac", ".aac"}
VIDEO_EXT = {".mp4", ".mov", ".webm", ".mkv", ".avi"}
OFFICE_EXT = {".pdf", ".docx", ".pptx", ".xlsx", ".csv"}
ALLOWED_EXT = TEXT_EXT | IMAGE_EXT | AUDIO_EXT | VIDEO_EXT | OFFICE_EXT

MIME_FALLBACK = {
    ".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".webp": "image/webp",
    ".gif": "image/gif", ".bmp": "image/bmp", ".mp3": "audio/mpeg", ".wav": "audio/wav",
    ".m4a": "audio/mp4", ".ogg": "audio/ogg", ".flac": "audio/flac", ".aac": "audio/aac",
    ".mp4": "video/mp4", ".mov": "video/quicktime", ".webm": "video/webm",
    ".mkv": "video/x-matroska", ".avi": "video/x-msvideo", ".pdf": "application/pdf",
}

MAGIC = {
    ".pdf": [b"%PDF"],
    ".docx": [b"PK\x03\x04"],
    ".pptx": [b"PK\x03\x04"],
    ".xlsx": [b"PK\x03\x04"],
    ".png": [b"\x89PNG"],
    ".jpg": [b"\xff\xd8\xff"],
    ".jpeg": [b"\xff\xd8\xff"],
    ".gif": [b"GIF8"],
}

ROWS_PER_SEGMENT = 40


def extension(filename: str) -> str:
    return PurePath(filename).suffix.lower()


def validate_upload(filename: str, head: bytes) -> str | None:
    """Return an error string, or None if the file looks acceptable."""
    ext = extension(filename)
    if ext not in ALLOWED_EXT:
        return f"Unsupported file type '{ext or filename}'"
    sigs = MAGIC.get(ext)
    if sigs and not any(head.startswith(s) for s in sigs):
        return "File content does not match its extension"
    return None


def guess_mime(filename: str) -> str:
    return MIME_FALLBACK.get(extension(filename), "application/octet-stream")


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16", "cp1252"):
        try:
            return data.decode(enc)
        except (UnicodeDecodeError, UnicodeError):
            continue
    return data.decode("utf-8", errors="replace")


def _rows_to_segments(sheet: str, header: list[str], rows: list[list[str]]) -> list[Segment]:
    segs = []
    for i in range(0, len(rows), ROWS_PER_SEGMENT):
        block = rows[i : i + ROWS_PER_SEGMENT]
        lines = [f"{sheet} | rows {i + 1}-{i + len(block)}"]
        for r in block:
            lines.append("; ".join(f"{h}: {v}" for h, v in zip(header, r) if v != ""))
        segs.append(Segment("\n".join(lines)))
    return segs


def extract_csv(data: bytes) -> list[Segment]:
    text = _decode(data)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = [r for r in csv.reader(io.StringIO(text), dialect) if any(c.strip() for c in r)]
    if not rows:
        return []
    return _rows_to_segments("CSV", rows[0], rows[1:])


def extract_xlsx(data: bytes) -> list[Segment]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    segs: list[Segment] = []
    for ws in wb.worksheets:
        rows = [["" if c is None else str(c) for c in r] for r in ws.iter_rows(values_only=True)]
        rows = [r for r in rows if any(c.strip() for c in r)]
        if rows:
            segs.extend(_rows_to_segments(f"Sheet {ws.title}", rows[0], rows[1:]))
    return segs


def extract_docx(data: bytes) -> list[Segment]:
    from docx import Document

    doc = Document(io.BytesIO(data))
    parts = []
    for p in doc.paragraphs:
        if not p.text.strip():
            continue
        level = p.style.name.lower().replace("heading ", "") if p.style and p.style.name.lower().startswith("heading") else None
        parts.append(f"{'#' * int(level)} {p.text}" if level and level.isdigit() else p.text)
    for t in doc.tables:
        for row in t.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return [Segment("\n\n".join(parts))] if parts else []


def extract_pptx(data: bytes) -> list[Segment]:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    segs = []
    for n, slide in enumerate(prs.slides, 1):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                texts.append(shape.text_frame.text.strip())
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    texts.append(" | ".join(c.text.strip() for c in row.cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            texts.append("Notes: " + slide.notes_slide.notes_text_frame.text.strip())
        if texts:
            segs.append(Segment(f"Slide {n}\n\n" + "\n\n".join(texts), page=n))
    return segs


OCR_PROMPT = (
    "Transcribe all text in this document exactly as written, preserving reading order, headings and tables "
    "(use markdown). Keep the original language. Then, if there are charts or images, describe them briefly."
)


async def extract_pdf(data: bytes, llm: LLMClient) -> list[Segment]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    segs = [Segment((p.extract_text() or "").strip(), page=i) for i, p in enumerate(reader.pages, 1)]
    total = sum(len(s.text) for s in segs)
    # Scanned PDF (little or no text layer): fall back to multimodal OCR.
    if total < 30 * max(1, len(segs)) // 2:
        text = await llm.describe_media(data, "application/pdf", OCR_PROMPT)
        return [Segment(text.strip())] if text.strip() else []
    return [s for s in segs if s.text]


async def extract(filename: str, data: bytes, llm: LLMClient) -> list[Segment]:
    ext = extension(filename)
    if ext in TEXT_EXT:
        t = _decode(data).strip()
        return [Segment(t)] if t else []
    if ext == ".csv":
        return extract_csv(data)
    if ext == ".xlsx":
        return extract_xlsx(data)
    if ext == ".docx":
        return extract_docx(data)
    if ext == ".pptx":
        return extract_pptx(data)
    if ext == ".pdf":
        return await extract_pdf(data, llm)
    mime = guess_mime(filename)
    if ext in IMAGE_EXT:
        text = await llm.describe_media(
            data, mime,
            "Extract any visible text verbatim (OCR), then describe the image in detail. "
            "Format: 'Text:' section followed by 'Description:' section.",
        )
        return [Segment(text.strip())] if text.strip() else []
    if ext in AUDIO_EXT or ext in VIDEO_EXT:
        kind = "video" if ext in VIDEO_EXT else "audio"
        text = await llm.describe_media(
            data, mime,
            f"Transcribe the speech in this {kind} verbatim in its original language, with approximate "
            f"timestamps like [mm:ss] every few sentences. After the transcript add a section "
            f"'Summary:' with a concise summary and key points." + (" Mention important visuals too." if kind == "video" else ""),
        )
        return [Segment(text.strip())] if text.strip() else []
    raise ValueError(f"Unsupported file type: {ext}")
