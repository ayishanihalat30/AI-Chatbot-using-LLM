import csv
import io

import pytest

from app.db import normalize_url
from app.services.chunking import Segment, chunk_segments
from app.services.extract import extract, validate_upload
from app.services.retrieval import Hit, rerank, rrf_fuse
from tests.conftest import FakeLLM


def test_chunking_respects_size_and_overlap():
    text = "\n\n".join(f"Paragraph {i}. " + "word " * 120 for i in range(20))
    chunks = chunk_segments([Segment(text, page=3)], max_chars=1500, overlap=200)
    assert len(chunks) > 3
    assert all(len(c.text) <= 1500 + 400 for c in chunks)
    assert all(c.page == 3 for c in chunks)
    assert chunks[1].text[:100].strip()  # overlap tail present


def test_chunking_splits_huge_unbroken_text_and_indic():
    chunks = chunk_segments([Segment("क" * 7000 + "। अगला वाक्य।")], max_chars=3000)
    assert len(chunks) >= 3


def test_chunking_empty():
    assert chunk_segments([Segment("   \n\n ")]) == []


def test_rrf_prefers_items_in_both_lists():
    assert rrf_fuse([["a", "b", "c"], ["c", "a", "d"]])[:2] == ["a", "c"]
    assert rrf_fuse([[], []]) == []


async def test_rerank_orders_by_llm_scores_and_falls_back():
    hits = [Hit(str(i), "d", "f", None, f"c{i}") for i in range(3)]

    class L(FakeLLM):
        async def generate(self, p):
            return "Scores: [1, 9, 4]"

    out = await rerank(L(), "q", hits, 2)
    assert [h.chunk_id for h in out] == ["1", "2"]
    out = await rerank(FakeLLM(), "q", hits, 2)  # unparsable -> original order
    assert [h.chunk_id for h in out] == ["0", "1"]


def test_validate_upload():
    assert validate_upload("a.exe", b"MZ") is not None
    assert validate_upload("a.pdf", b"not a pdf") is not None
    assert validate_upload("a.pdf", b"%PDF-1.7") is None
    assert validate_upload("notes.md", b"# hi") is None


def test_normalize_url_neon_style():
    url, args = normalize_url("postgresql://u:p@host/db?sslmode=require&channel_binding=require")
    assert url == "postgresql+asyncpg://u:p@host/db" and args == {"ssl": True}


async def test_extract_text_csv():
    llm = FakeLLM()
    segs = await extract("n.md", "# Title\n\nbody".encode(), llm)
    assert "Title" in segs[0].text
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["name", "city"])
    for i in range(100):
        w.writerow([f"n{i}", "Kochi"])
    segs = await extract("d.csv", buf.getvalue().encode(), llm)
    assert len(segs) == 3 and "name: n0; city: Kochi" in segs[0].text


async def test_extract_office_formats():
    from docx import Document
    from openpyxl import Workbook
    from pptx import Presentation

    llm = FakeLLM()
    d = Document()
    d.add_heading("Report", 1)
    d.add_paragraph("Hello docx world")
    b = io.BytesIO()
    d.save(b)
    assert "Hello docx world" in (await extract("r.docx", b.getvalue(), llm))[0].text

    wb = Workbook()
    wb.active.append(["item", "qty"])
    wb.active.append(["apple", 3])
    b = io.BytesIO()
    wb.save(b)
    assert "item: apple; qty: 3" in (await extract("s.xlsx", b.getvalue(), llm))[0].text

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Launch plan"
    b = io.BytesIO()
    prs.save(b)
    segs = await extract("p.pptx", b.getvalue(), llm)
    assert segs[0].page == 1 and "Launch plan" in segs[0].text


async def test_extract_media_uses_llm():
    llm = FakeLLM()
    for name in ("a.png", "a.mp3", "a.mp4"):
        segs = await extract(name, b"\x00data", llm)
        assert "zebra" in segs[0].text


async def test_extract_scanned_pdf_falls_back_to_ocr():
    from pypdf import PdfWriter

    w = PdfWriter()
    w.add_blank_page(200, 200)
    b = io.BytesIO()
    w.write(b)
    segs = await extract("scan.pdf", b.getvalue(), FakeLLM())
    assert "zebra" in segs[0].text


async def test_extract_unknown_raises():
    with pytest.raises(ValueError):
        await extract("x.bin", b"1", FakeLLM())
